"""Quarantined parsing for administrator-supplied review evidence."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from app.security.filter import blocked_findings

ALLOWED_UPLOAD_EXTENSIONS = {".csv", ".md", ".pdf", ".txt", ".xlsx"}
ALLOWED_UPLOAD_MEDIA_TYPES = {
    ".csv": {"text/csv", "application/csv", "text/plain"},
    ".md": {"text/markdown", "text/plain"},
    ".pdf": {"application/pdf"},
    ".txt": {"text/plain"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
}
FACT_UPDATE_COLUMNS = {"product_id", "field_name", "proposed_value"}
MAX_EXTRACTED_CHARACTERS = 200_000
MAX_PREVIEW_CHARACTERS = 4_000
MAX_FACT_UPDATES = 100


@dataclass(frozen=True, slots=True)
class StagedParseResult:
    preview: dict[str, Any]
    fact_updates: list[dict[str, Any]]
    findings: list[str]


def upload_media_type_allowed(filename: str, content_type: str | None) -> bool:
    extension = Path(filename).suffix.casefold()
    normalized = (content_type or "").partition(";")[0].strip().casefold()
    return normalized in ALLOWED_UPLOAD_MEDIA_TYPES.get(extension, set())


def _decode_text(data: bytes) -> str:
    if b"\x00" in data:
        raise ValueError("Text uploads cannot contain binary NUL bytes")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Text uploads must use UTF-8 encoding") from exc


def _limited(text: str) -> str:
    return text[:MAX_PREVIEW_CHARACTERS]


def _parse_json_value(value: str) -> Any:
    value = value.strip()
    if not value:
        raise ValueError("proposed_value cannot be empty")
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError:
        parsed = value
    if isinstance(parsed, (dict, list)):
        raise ValueError("proposed_value must be a scalar value")
    return parsed


def _parse_csv(data: bytes) -> tuple[str, dict[str, Any], list[dict[str, Any]]]:
    text = _decode_text(data)
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("CSV extracted text is too large")
    reader = csv.DictReader(io.StringIO(text))
    headers = [str(item) for item in (reader.fieldnames or [])]
    rows = list(reader)
    preview_rows = [
        {key: str(value or "")[:240] for key, value in row.items() if key is not None}
        for row in rows[:8]
    ]
    updates: list[dict[str, Any]] = []
    if FACT_UPDATE_COLUMNS.issubset(headers):
        if len(rows) > MAX_FACT_UPDATES:
            raise ValueError(f"Fact update CSV supports at most {MAX_FACT_UPDATES} rows")
        for source_row, row in enumerate(rows, start=2):
            product_id = str(row.get("product_id") or "").strip()
            field_name = str(row.get("field_name") or "").strip()
            if not product_id or not field_name:
                raise ValueError(f"Missing product_id or field_name at CSV row {source_row}")
            updates.append(
                {
                    "product_id": product_id,
                    "field_name": field_name,
                    "proposed_value": _parse_json_value(str(row.get("proposed_value") or "")),
                    "source_row": source_row,
                    "note": str(row.get("note") or "").strip()[:1000] or None,
                }
            )
    return text, {"headers": headers, "row_count": len(rows), "rows": preview_rows}, updates


def _parse_text(data: bytes) -> tuple[str, dict[str, Any]]:
    text = _decode_text(data)
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("Extracted text is too large")
    return text, {"excerpt": _limited(text), "character_count": len(text)}


def _parse_pdf(data: bytes) -> tuple[str, dict[str, Any]]:
    if not data.startswith(b"%PDF-"):
        raise ValueError("PDF content signature does not match its extension")
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("Encrypted PDF files are not accepted")
    if len(reader.pages) > 50:
        raise ValueError("PDF uploads support at most 50 pages")
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(pages)
    if not text.strip():
        raise ValueError("PDF uploads must contain extractable text for privacy scanning")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("PDF extracted text is too large")
    return text, {
        "page_count": len(pages),
        "excerpt": _limited(text),
        "character_count": len(text),
    }


def _parse_xlsx(data: bytes) -> tuple[str, dict[str, Any]]:
    if not data.startswith(b"PK"):
        raise ValueError("XLSX content signature does not match its extension")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > 2_000 or sum(item.file_size for item in members) > 25 * 1024 * 1024:
            raise ValueError("XLSX expanded content exceeds the safe preview limit")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    if len(workbook.sheetnames) > 20:
        workbook.close()
        raise ValueError("XLSX uploads support at most 20 sheets")
    sheet_previews: list[dict[str, Any]] = []
    text_parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row > 200 or sheet.max_column > 50:
                raise ValueError("XLSX sheets support at most 200 rows and 50 columns")
            rows: list[list[str]] = []
            for row_index, row in enumerate(
                sheet.iter_rows(max_row=200, max_col=50, values_only=True), start=1
            ):
                values = [str(value)[:500] if value is not None else "" for value in row]
                text_parts.append(f"{sheet.title} row {row_index}: " + " | ".join(values))
                if row_index <= 8:
                    rows.append(values)
            sheet_previews.append({"name": sheet.title, "rows": rows})
    finally:
        workbook.close()
    text = "\n".join(text_parts)
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("XLSX extracted text is too large")
    return text, {"sheets": sheet_previews, "sheet_count": len(sheet_previews)}


def parse_staged_source(filename: str, data: bytes) -> StagedParseResult:
    """Parse a bounded preview and detect sensitive text without exposing matched values."""

    extension = Path(filename).suffix.casefold()
    if extension not in ALLOWED_UPLOAD_EXTENSIONS:
        raise ValueError(
            "Unsupported upload type; allowed types are CSV, XLSX, PDF, Markdown, and TXT"
        )
    if extension == ".csv":
        text, preview, updates = _parse_csv(data)
    elif extension in {".md", ".txt"}:
        text, preview = _parse_text(data)
        updates = []
    elif extension == ".pdf":
        text, preview = _parse_pdf(data)
        updates = []
    else:
        text, preview = _parse_xlsx(data)
        updates = []
    findings = sorted(blocked_findings(text))
    return StagedParseResult(
        preview={} if findings else preview,
        fact_updates=[] if findings else updates,
        findings=findings,
    )
